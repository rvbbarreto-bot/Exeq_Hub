"""Testes importação Excel NfeProduct."""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from apps.accounts.models import Tenant
from apps.nfe.models import NfeProduct
from apps.nfe.product_import import (
    DATA_START_ROW,
    MODEL_VERSION,
    SHEET_PRODUCTS,
    STATUS_ATUALIZACAO,
    STATUS_NOVO,
    STATUS_PROCESSADO,
    STATUS_SEM_ALTERACAO,
    COLUMN_SPEC,
    ProductImportError,
    commit_import,
    generate_template_xlsx,
    validate_workbook,
)
from apps.nfe.services import create_product


@pytest.fixture
def import_tenant(db, settings):
    settings.NFE_ENABLED = True
    return Tenant.objects.create(
        slug="nfe-import",
        legal_name="Import Tenant",
        document="11222333000181",
        settings={"nfe_enabled": True},
    )


def _xlsx_bytes(rows: list[dict[str, str]]) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(generate_template_xlsx()))
    ws = wb[SHEET_PRODUCTS]
    for offset, row in enumerate(rows):
        line = DATA_START_ROW + offset
        for col_idx, (tech, _i, _r, _l) in enumerate(COLUMN_SPEC, start=1):
            ws.cell(row=line, column=col_idx, value=row.get(tech, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_generate_template_has_sheets():
    wb = openpyxl.load_workbook(BytesIO(generate_template_xlsx()))
    assert SHEET_PRODUCTS in wb.sheetnames
    assert "Instruções" in wb.sheetnames
    assert "Listas" in wb.sheetnames


def test_validate_new_product(import_tenant, settings):
    settings.NFE_ENABLED = True
    data = _xlsx_bytes(
        [
            {
                "codigo": "IMP-01",
                "descricao": "Produto import",
                "ncm": "21069090",
                "valor_unitario": "10,50",
            }
        ]
    )
    preview = validate_workbook(tenant=import_tenant, file_bytes=data, filename="t.xlsx")
    assert preview.summary["novos"] == 1
    assert preview.rows[0].status == STATUS_NOVO


def test_validate_duplicate_in_file(import_tenant, settings):
    settings.NFE_ENABLED = True
    row = {"codigo": "DUP", "descricao": "A", "ncm": "21069090"}
    preview = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes([row, row]),
        filename="t.xlsx",
    )
    assert preview.summary["erros"] == 2
    assert "duplicado" in preview.rows[0].messages[0].lower()


def test_validate_update_existing(import_tenant, settings):
    settings.NFE_ENABLED = True
    create_product(
        tenant=import_tenant,
        code="EX-1",
        description="Antiga",
        ncm="21069090",
        unit_price_cents=1000,
    )
    preview = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [
                {
                    "codigo": "EX-1",
                    "descricao": "Nova descrição",
                    "ncm": "21069090",
                    "valor_unitario": "20,00",
                }
            ]
        ),
        filename="t.xlsx",
    )
    assert preview.summary["atualizacoes"] == 1
    assert preview.rows[0].status == STATUS_ATUALIZACAO
    assert any(ch.field == "description" for ch in preview.rows[0].changes)


def test_empty_cell_keeps_value_on_update(import_tenant, settings):
    settings.NFE_ENABLED = True
    create_product(
        tenant=import_tenant,
        code="KEEP",
        description="Desc original",
        ncm="21069090",
        unit_price_cents=5000,
    )
    preview = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [{"codigo": "KEEP", "descricao": "Desc original", "ncm": "21069090"}]
        ),
        filename="t.xlsx",
    )
    assert preview.summary["sem_alteracao"] == 1


def test_decimal_formats(import_tenant, settings):
    settings.NFE_ENABLED = True
    for price in ("10", "10,5", "10.50"):
        preview = validate_workbook(
            tenant=import_tenant,
            file_bytes=_xlsx_bytes(
                [
                    {
                        "codigo": f"P-{price}",
                        "descricao": "P",
                        "ncm": "21069090",
                        "valor_unitario": price,
                    }
                ]
            ),
            filename="t.xlsx",
        )
        assert preview.rows[0].parsed["unit_price_cents"] >= 1000


def test_commit_create_and_update(import_tenant, settings):
    settings.NFE_ENABLED = True
    preview = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [
                {
                    "codigo": "NEW-1",
                    "descricao": "Novo",
                    "ncm": "21069090",
                    "valor_unitario": "15,00",
                }
            ]
        ),
        filename="t.xlsx",
    )
    result = commit_import(tenant=import_tenant, token=preview.token, actor="test")
    assert result.summary["cadastrados"] == 1
    prod = NfeProduct.objects.get(tenant=import_tenant, code="NEW-1")
    assert prod.unit_price_cents == 1500

    preview2 = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [
                {
                    "codigo": "NEW-1",
                    "descricao": "Novo",
                    "ncm": "21069090",
                    "valor_unitario": "18,00",
                }
            ]
        ),
        filename="t.xlsx",
    )
    result2 = commit_import(tenant=import_tenant, token=preview2.token, actor="test")
    assert result2.summary["alterados"] == 1
    prod.refresh_from_db()
    assert prod.unit_price_cents == 1800


def test_invalid_header_rejected(import_tenant, settings):
    settings.NFE_ENABLED = True
    wb = openpyxl.load_workbook(BytesIO(generate_template_xlsx()))
    ws = wb[SHEET_PRODUCTS]
    ws.cell(row=1, column=1, value="codigo_errado")
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ProductImportError, match="obrigatórias ausentes|não reconhecida"):
        validate_workbook(tenant=import_tenant, file_bytes=buf.getvalue(), filename="t.xlsx")


def test_inactive_not_reactivated_without_explicit(import_tenant, settings):
    settings.NFE_ENABLED = True
    p = create_product(
        tenant=import_tenant,
        code="INACT",
        description="Off",
        ncm="21069090",
        is_active=False,
    )
    preview = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [{"codigo": "INACT", "descricao": "Off", "ncm": "21069090"}]
        ),
        filename="t.xlsx",
    )
    assert preview.rows[0].parsed["is_active"] is False

    preview2 = validate_workbook(
        tenant=import_tenant,
        file_bytes=_xlsx_bytes(
            [
                {
                    "codigo": "INACT",
                    "descricao": "Off",
                    "ncm": "21069090",
                    "ativo": "Sim",
                }
            ]
        ),
        filename="t.xlsx",
    )
    assert preview2.rows[0].parsed["is_active"] is True
