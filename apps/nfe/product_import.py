"""Importação/atualização em massa NfeProduct via Excel (.xlsx) — modelo v1.0."""

from __future__ import annotations

import json
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from django.conf import settings
from django.db import transaction

from apps.fiscal.goods_catalog import dropdown_units
from apps.nfe.exceptions import NfeValidationError
from apps.nfe.models import NfeProduct
from apps.nfe.services import create_product, update_product

MODEL_VERSION = "1.0"
SHEET_PRODUCTS = "Produtos"
SHEET_INSTRUCTIONS = "Instruções"
SHEET_LISTS = "Listas"
TEMPLATE_FILENAME = "modelo_importacao_produtos_nfe.xlsx"
HEADER_ROW = 1
LABEL_ROW = 2
DATA_START_ROW = 3

FORMULA_PREFIXES = ("=", "+", "-", "@")

# cabeçalho técnico linha 1 → campo interno
COLUMN_SPEC: tuple[tuple[str, str, bool, str], ...] = (
    ("codigo", "code", True, "Código do Produto — OBRIGATÓRIO — TEXTO"),
    ("descricao", "description", True, "Descrição — OBRIGATÓRIO — TEXTO"),
    ("ncm", "ncm", True, "NCM — OBRIGATÓRIO — 8 DÍGITOS"),
    (
        "valor_unitario",
        "unit_price",
        False,
        "Valor Unitário (R$) — OPCIONAL — DECIMAL (ex.: 10,50)",
    ),
    ("unidade", "unit", False, "Unidade — OPCIONAL — LISTA"),
    ("origem", "origin", False, "Origem — OPCIONAL — 0 a 8"),
    ("cfop_interno", "cfop_internal", False, "CFOP Interno — OPCIONAL — 4 DÍGITOS"),
    (
        "cfop_interestadual",
        "cfop_interstate",
        False,
        "CFOP Interestadual — OPCIONAL — 4 DÍGITOS",
    ),
    ("csosn", "csosn", False, "CSOSN — OPCIONAL — 3 DÍGITOS"),
    ("cst_icms", "icms_cst", False, "CST ICMS — OPCIONAL — 3 DÍGITOS"),
    (
        "aliquota_icms",
        "icms_rate",
        False,
        "Alíquota ICMS (%) — OPCIONAL — DECIMAL/PERCENTUAL",
    ),
    ("cst_pis", "pis_cst", False, "CST PIS — OPCIONAL — 2 DÍGITOS"),
    (
        "aliquota_pis",
        "pis_rate",
        False,
        "Alíquota PIS (%) — OPCIONAL — DECIMAL/PERCENTUAL",
    ),
    ("cst_cofins", "cofins_cst", False, "CST COFINS — OPCIONAL — 2 DÍGITOS"),
    (
        "aliquota_cofins",
        "cofins_rate",
        False,
        "Alíquota COFINS (%) — OPCIONAL — DECIMAL/PERCENTUAL",
    ),
    ("gtin", "gtin", False, "GTIN/EAN — OPCIONAL — TEXTO (preservar zeros)"),
    ("ativo", "is_active", False, "Ativo — OPCIONAL — Sim/Não"),
)

TECH_HEADERS = [spec[0] for spec in COLUMN_SPEC]
REQUIRED_HEADERS = {spec[0] for spec in COLUMN_SPEC if spec[2]}
INTERNAL_FIELDS = {spec[1] for spec in COLUMN_SPEC}

STATUS_NOVO = "NOVO"
STATUS_ATUALIZACAO = "ATUALIZAÇÃO"
STATUS_SEM_ALTERACAO = "SEM ALTERAÇÃO"
STATUS_ERRO = "ERRO"
STATUS_NAO_PROCESSADO = "NÃO PROCESSADO"
STATUS_PROCESSADO = "PROCESSADO"

ORIGIN_LABELS = {
    "0": "0 — Nacional",
    "1": "1 — Estrangeira — Importação direta",
    "2": "2 — Estrangeira — Adquirida no mercado interno",
    "3": "3 — Nacional — Conteúdo importação 40–70%",
    "4": "4 — Nacional — Processos produtivos básicos",
    "5": "5 — Nacional — Conteúdo importação ≤ 40%",
    "6": "6 — Estrangeira — Importação direta sem similar",
    "7": "7 — Estrangeira — Mercado interno sem similar",
    "8": "8 — Nacional — Conteúdo importação > 70%",
}


class ProductImportError(ValueError):
    pass


@dataclass
class FieldChange:
    field: str
    label: str
    old: str
    new: str


@dataclass
class ImportRowPreview:
    line_number: int
    code: str
    operation: str
    status: str
    messages: list[str] = field(default_factory=list)
    changes: list[FieldChange] = field(default_factory=list)
    parsed: dict[str, Any] = field(default_factory=dict)
    product_id: str | None = None


@dataclass
class ImportPreview:
    token: str
    tenant_id: str
    filename: str
    model_version: str
    summary: dict[str, int]
    rows: list[ImportRowPreview]
    created_at: str


def max_rows() -> int:
    return int(getattr(settings, "NFE_PRODUCT_IMPORT_MAX_ROWS", 500) or 500)


def max_bytes() -> int:
    return int(
        getattr(settings, "NFE_PRODUCT_IMPORT_MAX_BYTES", 3 * 1024 * 1024)
        or (3 * 1024 * 1024)
    )


def _storage_root() -> Path:
    root = Path(getattr(settings, "LOCAL_STORAGE_ROOT", settings.BASE_DIR / ".storage"))
    return root / "nfe_product_import"


def preview_path(tenant_id: str, token: str) -> Path:
    return _storage_root() / str(tenant_id) / f"{token}.json"


def report_path(tenant_id: str, token: str) -> Path:
    return _storage_root() / str(tenant_id) / f"{token}-report.json"


def _load_openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover
        raise ProductImportError(
            "Dependência openpyxl ausente. Instale com: pip install openpyxl"
        ) from exc
    return openpyxl, Alignment, Font, PatternFill, get_column_letter, DataValidation


def _norm_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    text = str(value).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        return text[:-2]
    return text


def sanitize_text(value: str) -> str:
    text = (value or "").strip()
    if text and text[0] in FORMULA_PREFIXES:
        return f"'{text}"
    return text


def _parse_brl_to_cents(raw: str) -> int:
    from apps.hub_v4.forms import parse_brl_amount_cents

    return parse_brl_amount_cents(raw, field_label="Valor unitário", allow_zero=True)


def _parse_percent_to_bp(raw: str) -> int:
    from apps.hub_v4.forms import _parse_percent_to_bp

    return _parse_percent_to_bp(raw)


def _parse_is_active(raw: str) -> bool | None:
    text = (raw or "").strip().lower()
    if not text:
        return None
    if text in {"1", "sim", "s", "true", "yes", "y", "ativo"}:
        return True
    if text in {"0", "não", "nao", "n", "false", "no", "inativo"}:
        return False
    raise ProductImportError(f"Valor de ativo inválido: {raw!r}. Use Sim ou Não.")


def _field_labels() -> dict[str, str]:
    return {spec[1]: spec[3].split(" — ")[0] for spec in COLUMN_SPEC}


def _display_value(field: str, value: Any) -> str:
    if field == "unit_price_cents":
        cents = int(value or 0)
        return f"R$ {cents / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if field.endswith("_rate_bp"):
        bp = int(value or 0)
        return f"{bp / 100:.2f}%".replace(".", ",")
    if field == "is_active":
        return "Sim" if value else "Não"
    return str(value if value is not None else "")


def generate_template_xlsx() -> bytes:
    openpyxl, Alignment, Font, PatternFill, get_column_letter, DataValidation = _load_openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_PRODUCTS

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    label_fill = PatternFill("solid", fgColor="E8EEF4")

    for col_idx, (tech, _internal, _req, label) in enumerate(COLUMN_SPEC, start=1):
        c1 = ws.cell(row=HEADER_ROW, column=col_idx, value=tech)
        c1.fill = header_fill
        c1.font = header_font
        c1.alignment = Alignment(horizontal="center")
        c2 = ws.cell(row=LABEL_ROW, column=col_idx, value=label)
        c2.fill = label_fill
        c2.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(label) // 2)

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMN_SPEC))}{HEADER_ROW}"

    for col_idx, tech in enumerate(TECH_HEADERS, start=1):
        if tech in {"codigo", "ncm", "gtin", "csosn", "cst_icms", "cst_pis", "cst_cofins"}:
            for row_idx in range(DATA_START_ROW, DATA_START_ROW + 200):
                ws.cell(row=row_idx, column=col_idx).number_format = "@"

    inst = wb.create_sheet(SHEET_INSTRUCTIONS)
    instructions = [
        ("Finalidade", "Importar ou atualizar produtos NF-e em massa (modelo v1.0)."),
        ("Versão do modelo", MODEL_VERSION),
        ("Identificação", "Chave única por empresa: codigo (código do produto)."),
        ("Novos produtos", "codigo + descricao + ncm obrigatórios na linha."),
        ("Atualização", "Célula vazia mantém o valor atual no banco."),
        ("codigo", "Não pode ser alterado após cadastro (use novo código para novo SKU)."),
        ("valor_unitario", "Decimal BR: 10,50 ou 1500,00. Aceita zero."),
        ("Alíquotas %", "Informe percentual (ex.: 18 ou 18,5). Armazenamento interno em basis points."),
        ("gtin", "Texto — preserve zeros à esquerda. Vazio = sem GTIN."),
        ("ativo", "Sim/Não. Reativação só com Sim explícito. Vazio mantém status atual."),
        ("Limite", f"Máximo {max_rows()} linhas e {max_bytes() // (1024 * 1024)} MB por arquivo."),
        ("Exclusão", "Produtos nunca são excluídos pela importação."),
    ]
    inst["A1"] = "Campo"
    inst["B1"] = "Instrução"
    inst["A1"].font = Font(bold=True)
    inst["B1"].font = Font(bold=True)
    for idx, (title, text) in enumerate(instructions, start=2):
        inst.cell(row=idx, column=1, value=title)
        cell = inst.cell(row=idx, column=2, value=text)
        if title == "Versão do modelo":
            cell.number_format = "@"
    inst.column_dimensions["A"].width = 22
    inst.column_dimensions["B"].width = 80

    lists = wb.create_sheet(SHEET_LISTS)
    lists.sheet_state = "hidden"
    lists["A1"] = "origem"
    lists["B1"] = "unidade"
    lists["C1"] = "ativo"
    for idx, code in enumerate(sorted(ORIGIN_LABELS), start=2):
        lists.cell(row=idx, column=1, value=code)
    units = [code for code, _ in dropdown_units()]
    for idx, code in enumerate(units, start=2):
        lists.cell(row=idx, column=2, value=code)
    for idx, val in enumerate(["Sim", "Não"], start=2):
        lists.cell(row=idx, column=3, value=val)

    unit_col = TECH_HEADERS.index("unidade") + 1
    unit_letter = get_column_letter(unit_col)
    dv_unit = DataValidation(
        type="list",
        formula1=f"={SHEET_LISTS}!$B$2:$B${1 + len(units)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_unit)
    dv_unit.add(f"{unit_letter}{DATA_START_ROW}:{unit_letter}504")

    ativo_col = TECH_HEADERS.index("ativo") + 1
    ativo_letter = get_column_letter(ativo_col)
    dv_ativo = DataValidation(
        type="list",
        formula1=f"={SHEET_LISTS}!$C$2:$C$3",
        allow_blank=True,
    )
    ws.add_data_validation(dv_ativo)
    dv_ativo.add(f"{ativo_letter}{DATA_START_ROW}:{ativo_letter}504")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_header_map(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col_idx in range(1, len(TECH_HEADERS) + 1):
        key = _norm_cell(ws.cell(row=HEADER_ROW, column=col_idx).value).lower()
        if key:
            headers[key] = col_idx
    return headers


def _validate_structure(ws, wb) -> str:
    header_map = _read_header_map(ws)
    missing = REQUIRED_HEADERS - set(header_map)
    if missing:
        raise ProductImportError(
            f"Colunas obrigatórias ausentes: {', '.join(sorted(missing))}. "
            "Utilize o modelo oficial disponibilizado pelo sistema."
        )
    unknown = set(header_map) - set(TECH_HEADERS)
    if unknown:
        raise ProductImportError(
            f"Coluna(s) não reconhecida(s): {', '.join(sorted(unknown))}. "
            "Utilize o modelo oficial disponibilizado pelo sistema."
        )
    version = MODEL_VERSION
    if SHEET_INSTRUCTIONS in wb.sheetnames:
        inst = wb[SHEET_INSTRUCTIONS]
        for row in inst.iter_rows(min_row=2, max_col=2, values_only=True):
            if row and str(row[0] or "").strip().lower() == "versão do modelo":
                version = _norm_cell(row[1]) or version
                break
    if version not in {MODEL_VERSION, "1", "1.0"}:
        raise ProductImportError(
            f"Versão do modelo {version!r} incompatível. Esperado {MODEL_VERSION}."
        )
    return MODEL_VERSION


def _row_is_empty(ws, row_idx: int, header_map: dict[str, int]) -> bool:
    for key in REQUIRED_HEADERS:
        col = header_map.get(key)
        if col and _norm_cell(ws.cell(row=row_idx, column=col).value):
            return False
    for key, col in header_map.items():
        if key in REQUIRED_HEADERS:
            continue
        if _norm_cell(ws.cell(row=row_idx, column=col).value):
            return False
    return True


def _parse_row_cells(
    ws,
    row_idx: int,
    header_map: dict[str, int],
) -> tuple[dict[str, str], set[str]]:
    """Retorna valores raw por col técnica e conjunto de colunas presentes (não vazias)."""
    raw: dict[str, str] = {}
    provided: set[str] = set()
    for tech, internal, _req, _label in COLUMN_SPEC:
        col = header_map.get(tech)
        if not col:
            continue
        text = sanitize_text(_norm_cell(ws.cell(row=row_idx, column=col).value))
        raw[tech] = text
        if text != "":
            provided.add(internal)
    return raw, provided


def _build_payload(
    raw: dict[str, str],
    provided: set[str],
    existing: NfeProduct | None,
) -> dict[str, Any]:
    errors: list[str] = []
    is_create = existing is None
    out: dict[str, Any] = (
        {
            "unit": "UN",
            "origin": "0",
            "cfop_internal": "5102",
            "cfop_interstate": "6102",
            "csosn": "102",
            "icms_cst": "",
            "icms_rate_bp": 0,
            "pis_cst": "07",
            "pis_rate_bp": 0,
            "cofins_cst": "07",
            "cofins_rate_bp": 0,
            "gtin": "",
            "unit_price_cents": 0,
            "is_active": True,
        }
        if is_create
        else _product_snapshot(existing)
    )

    code = (raw.get("codigo") or "").strip()[:60]
    if not code:
        errors.append("Código do produto é obrigatório.")
    out["code"] = code

    if "description" in provided or is_create:
        desc = (raw.get("descricao") or "").strip()[:120]
        if not desc:
            errors.append("Descrição é obrigatória.")
        else:
            out["description"] = desc

    if "ncm" in provided or is_create:
        ncm = "".join(ch for ch in (raw.get("ncm") or "") if ch.isdigit())[:8]
        if len(ncm) != 8:
            errors.append("NCM deve ter 8 dígitos.")
        else:
            out["ncm"] = ncm

    if "unit_price" in provided:
        try:
            out["unit_price_cents"] = _parse_brl_to_cents(raw.get("valor_unitario") or "")
        except ValueError as exc:
            errors.append(str(exc))

    for tech, internal in (
        ("unidade", "unit"),
        ("origem", "origin"),
        ("cfop_interno", "cfop_internal"),
        ("cfop_interestadual", "cfop_interstate"),
        ("csosn", "csosn"),
        ("cst_icms", "icms_cst"),
        ("cst_pis", "pis_cst"),
        ("cst_cofins", "cofins_cst"),
        ("gtin", "gtin"),
    ):
        if internal in provided:
            val = (raw.get(tech) or "").strip()
            out[internal] = val[:6] if internal == "unit" else val

    if "icms_rate" in provided:
        try:
            out["icms_rate_bp"] = _parse_percent_to_bp(raw.get("aliquota_icms") or "0")
        except ValueError as exc:
            errors.append(str(exc))
    if "pis_rate" in provided:
        try:
            out["pis_rate_bp"] = _parse_percent_to_bp(raw.get("aliquota_pis") or "0")
        except ValueError as exc:
            errors.append(str(exc))
    if "cofins_rate" in provided:
        try:
            out["cofins_rate_bp"] = _parse_percent_to_bp(raw.get("aliquota_cofins") or "0")
        except ValueError as exc:
            errors.append(str(exc))

    if "is_active" in provided:
        try:
            parsed = _parse_is_active(raw.get("ativo") or "")
            if parsed is None:
                errors.append("Informe Sim ou Não na coluna ativo.")
            else:
                out["is_active"] = parsed
        except ProductImportError as exc:
            errors.append(str(exc))
    elif is_create:
        out["is_active"] = True

    if errors:
        raise ProductImportError("; ".join(errors))
    return out


def _product_snapshot(product: NfeProduct) -> dict[str, Any]:
    return {
        "code": product.code,
        "description": product.description,
        "ncm": product.ncm,
        "unit_price_cents": product.unit_price_cents,
        "unit": product.unit,
        "origin": product.origin,
        "cfop_internal": product.cfop_internal,
        "cfop_interstate": product.cfop_interstate,
        "csosn": product.csosn,
        "icms_cst": product.icms_cst,
        "icms_rate_bp": product.icms_rate_bp,
        "pis_cst": product.pis_cst,
        "pis_rate_bp": product.pis_rate_bp,
        "cofins_cst": product.cofins_cst,
        "cofins_rate_bp": product.cofins_rate_bp,
        "gtin": product.gtin,
        "is_active": product.is_active,
    }


def _diff_product(
    before: dict[str, Any],
    after: dict[str, Any],
) -> list[FieldChange]:
    labels = {
        "description": "Descrição",
        "ncm": "NCM",
        "unit_price_cents": "Preço unitário",
        "unit": "Unidade",
        "origin": "Origem",
        "cfop_internal": "CFOP interno",
        "cfop_interstate": "CFOP interestadual",
        "csosn": "CSOSN",
        "icms_cst": "CST ICMS",
        "icms_rate_bp": "Alíquota ICMS",
        "pis_cst": "CST PIS",
        "pis_rate_bp": "Alíquota PIS",
        "cofins_cst": "CST COFINS",
        "cofins_rate_bp": "Alíquota COFINS",
        "gtin": "GTIN/EAN",
        "is_active": "Ativo",
    }
    changes: list[FieldChange] = []
    for key, label in labels.items():
        if before.get(key) != after.get(key):
            changes.append(
                FieldChange(
                    field=key,
                    label=label,
                    old=_display_value(key, before.get(key)),
                    new=_display_value(key, after.get(key)),
                )
            )
    return changes


def validate_workbook(*, tenant, file_bytes: bytes, filename: str) -> ImportPreview:
    if len(file_bytes) > max_bytes():
        raise ProductImportError(
            f"Arquivo excede o limite de {max_bytes() // (1024 * 1024)} MB."
        )
    if not (filename or "").lower().endswith(".xlsx"):
        raise ProductImportError("Envie um arquivo Excel (.xlsx).")

    openpyxl, *_rest = _load_openpyxl()
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ProductImportError("Arquivo Excel inválido ou corrompido.") from exc

    if SHEET_PRODUCTS not in wb.sheetnames:
        wb.close()
        raise ProductImportError(f"Aba '{SHEET_PRODUCTS}' não encontrada.")
    ws = wb[SHEET_PRODUCTS]
    model_version = _validate_structure(ws, wb)

    header_map = _read_header_map(ws)
    data_rows: list[tuple[int, dict[str, str], set[str]]] = []
    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        if _row_is_empty(ws, row_idx, header_map):
            continue
        raw, provided = _parse_row_cells(ws, row_idx, header_map)
        data_rows.append((row_idx, raw, provided))
    wb.close()

    if not data_rows:
        raise ProductImportError("Nenhuma linha de produto encontrada na planilha.")
    if len(data_rows) > max_rows():
        raise ProductImportError(f"Máximo de {max_rows()} produtos por arquivo.")

    codes = [(raw.get("codigo") or "").strip()[:60] for _, raw, _ in data_rows]
    dup_in_file: dict[str, list[int]] = {}
    for (line, raw, _), code in zip(data_rows, codes, strict=True):
        if not code:
            continue
        dup_in_file.setdefault(code, []).append(line)

    existing_map = {
        p.code: p
        for p in NfeProduct.objects.filter(tenant=tenant, code__in=[c for c in codes if c])
    }

    previews: list[ImportRowPreview] = []
    summary = {
        "total": 0,
        "novos": 0,
        "atualizacoes": 0,
        "sem_alteracao": 0,
        "erros": 0,
    }

    for line_number, raw, provided in data_rows:
        code = (raw.get("codigo") or "").strip()[:60]
        row_preview = ImportRowPreview(
            line_number=line_number,
            code=code or "—",
            operation=STATUS_NOVO,
            status=STATUS_ERRO,
            messages=[],
            changes=[],
            parsed={},
        )
        summary["total"] += 1

        if code and len(dup_in_file.get(code, [])) > 1:
            row_preview.status = STATUS_NAO_PROCESSADO
            row_preview.messages.append(
                f"Produto duplicado dentro da planilha (linhas: {', '.join(map(str, dup_in_file[code]))})."
            )
            summary["erros"] += 1
            previews.append(row_preview)
            continue

        existing = existing_map.get(code)
        is_create = existing is None

        try:
            payload = _build_payload(raw, provided, existing)
        except ProductImportError as exc:
            row_preview.messages.append(str(exc))
            summary["erros"] += 1
            previews.append(row_preview)
            continue

        if is_create:
            row_preview.operation = STATUS_NOVO
            row_preview.status = STATUS_NOVO
            row_preview.parsed = payload
            summary["novos"] += 1
        else:
            row_preview.product_id = str(existing.id)
            before = _product_snapshot(existing)
            after = payload
            if before.get("code") != after.get("code"):
                row_preview.status = STATUS_NAO_PROCESSADO
                row_preview.messages.append(
                    "O campo código não pode ser alterado após o cadastro."
                )
                summary["erros"] += 1
                previews.append(row_preview)
                continue
            changes = _diff_product(before, after)
            row_preview.parsed = payload
            if changes:
                row_preview.operation = STATUS_ATUALIZACAO
                row_preview.status = STATUS_ATUALIZACAO
                row_preview.changes = changes
                summary["atualizacoes"] += 1
            else:
                row_preview.operation = STATUS_SEM_ALTERACAO
                row_preview.status = STATUS_SEM_ALTERACAO
                summary["sem_alteracao"] += 1
        previews.append(row_preview)

    token = uuid.uuid4().hex
    preview = ImportPreview(
        token=token,
        tenant_id=str(tenant.id),
        filename=filename,
        model_version=model_version,
        summary=summary,
        rows=previews,
        created_at=datetime.now(timezone.utc).isoformat(),
    )
    save_preview(preview)
    return preview


def save_preview(preview: ImportPreview) -> None:
    path = preview_path(preview.tenant_id, preview.token)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "token": preview.token,
        "tenant_id": preview.tenant_id,
        "filename": preview.filename,
        "model_version": preview.model_version,
        "summary": preview.summary,
        "created_at": preview.created_at,
        "rows": [
            {
                **asdict(r),
                "changes": [asdict(c) for c in r.changes],
            }
            for r in preview.rows
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def load_preview(tenant_id: str, token: str) -> ImportPreview:
    path = preview_path(tenant_id, token)
    if not path.is_file():
        raise ProductImportError("Prévia expirada ou inválida. Envie o arquivo novamente.")
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = []
    for raw in data.get("rows") or []:
        changes = [FieldChange(**c) for c in raw.pop("changes", [])]
        rows.append(ImportRowPreview(changes=changes, **raw))
    return ImportPreview(rows=rows, **{k: data[k] for k in data if k != "rows"})


def commit_import(*, tenant, token: str, actor: str) -> ImportPreview:
    preview = load_preview(str(tenant.id), token)
    if preview.tenant_id != str(tenant.id):
        raise ProductImportError("Prévia não pertence a este tenant.")

    existing_map = {
        p.code: p
        for p in NfeProduct.objects.filter(
            tenant=tenant,
            code__in=[r.code for r in preview.rows if r.code and r.code != "—"],
        )
    }

    result_summary = {
        "total": preview.summary.get("total", 0),
        "cadastrados": 0,
        "alterados": 0,
        "sem_alteracao": 0,
        "nao_processados": 0,
    }

    for row in preview.rows:
        if row.status in {STATUS_ERRO, STATUS_NAO_PROCESSADO}:
            result_summary["nao_processados"] += 1
            continue
        if row.status == STATUS_SEM_ALTERACAO:
            row.status = STATUS_PROCESSADO
            result_summary["sem_alteracao"] += 1
            continue

        payload = row.parsed or {}
        code = payload.get("code") or row.code
        try:
            with transaction.atomic():
                if row.operation == STATUS_NOVO:
                    create_product(tenant=tenant, **payload)
                    row.status = STATUS_PROCESSADO
                    result_summary["cadastrados"] += 1
                elif row.operation == STATUS_ATUALIZACAO:
                    product = existing_map.get(code)
                    if product is None:
                        raise ProductImportError(f"Produto {code} não encontrado.")
                    update_kwargs = {
                        k: v
                        for k, v in payload.items()
                        if k != "code"
                    }
                    update_product(product, **update_kwargs)
                    row.status = STATUS_PROCESSADO
                    result_summary["alterados"] += 1
        except (NfeValidationError, ProductImportError, ValueError) as exc:
            row.status = STATUS_NAO_PROCESSADO
            row.messages = [str(exc)]
            result_summary["nao_processados"] += 1
        except Exception as exc:
            row.status = STATUS_NAO_PROCESSADO
            row.messages = ["Não foi possível processar este produto. Verifique os dados."]
            result_summary["nao_processados"] += 1
            _ = exc

    preview.summary = result_summary
    save_preview(preview)
    report_file = report_path(str(tenant.id), token)
    report_file.write_text(
        json.dumps(
            {
                "token": token,
                "actor": actor,
                "filename": preview.filename,
                "summary": result_summary,
                "rows": [
                    {
                        "line_number": r.line_number,
                        "code": r.code,
                        "operation": r.operation,
                        "status": r.status,
                        "messages": r.messages,
                        "changes": [asdict(c) for c in r.changes],
                    }
                    for r in preview.rows
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return preview


def generate_report_xlsx(preview: ImportPreview) -> bytes:
    openpyxl, Alignment, Font, PatternFill, get_column_letter, _dv = _load_openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório"
    headers = [
        "Linha",
        "Produto",
        "Operação",
        "Status",
        "Campo",
        "Valor atual",
        "Valor informado",
        "Motivo",
    ]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
    row_idx = 2
    for item in preview.rows:
        if item.changes:
            for ch in item.changes:
                ws.cell(row=row_idx, column=1, value=item.line_number)
                ws.cell(row=row_idx, column=2, value=item.code)
                ws.cell(row=row_idx, column=3, value=item.operation)
                ws.cell(row=row_idx, column=4, value=item.status)
                ws.cell(row=row_idx, column=5, value=ch.label)
                ws.cell(row=row_idx, column=6, value=ch.old)
                ws.cell(row=row_idx, column=7, value=ch.new)
                ws.cell(row=row_idx, column=8, value="; ".join(item.messages))
                row_idx += 1
        else:
            ws.cell(row=row_idx, column=1, value=item.line_number)
            ws.cell(row=row_idx, column=2, value=item.code)
            ws.cell(row=row_idx, column=3, value=item.operation)
            ws.cell(row=row_idx, column=4, value=item.status)
            ws.cell(row=row_idx, column=8, value="; ".join(item.messages))
            row_idx += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
